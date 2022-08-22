import openpyxl
from openpyxl.chart import StockChart, Reference, LineChart, BarChart
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.updown_bars import UpDownBars
from openpyxl.drawing.line import LineProperties

# 绘制StockChart
def handle_stock_chart(src, des):
    wb = openpyxl.load_workbook(src)

    # 包含数据的工作表
    data_sheet = wb['000001']

    # 图形工作表
    chart_sheet = wb.create_sheet('图形')
    wb.active = chart_sheet

    sc = StockChart()

    sc.width = 35
    sc.height = 15

    max_row_number = 50
    # 数据选区
    data = Reference(data_sheet, min_col=2, max_col=5, min_row=1, max_row=max_row_number)
    sc.add_data(data, titles_from_data=True)

    labels = Reference(data_sheet, min_col=1, min_row=2, max_row=max_row_number)
    sc.set_categories(labels)

    for c in sc.series:
        c.graphicalProperties.line.noFill = True

    sc.hiLowLines = ChartLines()
    sc.upDownBars = UpDownBars()

    # 由于Excel中的BUG错误，仅当数据序列中的至少一个具有一些虚拟值时，才会显示高/低行
    from openpyxl.chart.data_source import NumData, NumVal
    pts = [NumVal(idx=i) for i in range(len(data) - 1)]
    cache = NumData(pt=pts)
    sc.series[-1].val.numRef.numCache = cache

    close_line = LineChart()
    line_data = Reference(data_sheet, min_col=5, min_row=1, max_row=max_row_number)
    close_line.add_data(line_data, titles_from_data=True)

    s1 = close_line.series[0]
    lineproperties = LineProperties(solidFill='FF0000', w=1)
    s1.graphicalProperties.line = lineproperties

    sc += close_line

    bar = BarChart()
    bar_data = Reference(data_sheet, min_col=6, min_row=1, max_row=max_row_number)
    bar.add_data(bar_data, titles_from_data=True)
    bar.y_axis.axId = 200
    bar.y_axis.majorGridlines = None
    bar.y_axis.crosses = 'max'
    bar.y_axis.title = '成交量'

    sc += bar

    sc.title = '上证指数'

    chart_sheet.add_chart(sc, anchor='B2')

    wb.save(des)


if __name__ == "__main__":
    handle_stock_chart('./000001.xlsx', './pExcel_08.xlsx')